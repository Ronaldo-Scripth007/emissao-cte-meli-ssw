import os
import sys
import time
import csv
import unicodedata
from datetime import datetime
from pathlib import Path
import pandas as pd
from dotenv import load_dotenv
from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC


def obter_pasta_aplicacao() -> Path:
    """
    Define a pasta base da automação.

    Executando como arquivo .py:
        utiliza a pasta onde está o script.

    Executando como arquivo .exe dentro de uma subpasta:
        utiliza a pasta principal, um nível acima da pasta do executável.
    """
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent.parent

    return Path(__file__).resolve().parent


SCRIPT_DIR = obter_pasta_aplicacao()

print(f"[info] Pasta da aplicação: {SCRIPT_DIR}")

# =====================================================================
# CONFIGURAÇÕES
# =====================================================================
ENV_PATH = SCRIPT_DIR / ".env"

if not ENV_PATH.exists():
    raise FileNotFoundError(f"Arquivo .env não encontrado na pasta da aplicação:\n{ENV_PATH}")

load_dotenv(dotenv_path=ENV_PATH)

SSW_URL = os.getenv("SSW_URL", "").strip()
SSW_DOMINIO = os.getenv("SSW_DOMINIO", "").strip()
SSW_CPF = os.getenv("SSW_CPF", "").strip()
SSW_USER = os.getenv("SSW_USER", "").strip()
SSW_PASS = os.getenv("SSW_PASS", "").strip()
SSW_TIPO_MERCADORIA = os.getenv("SSW_TIPO_MERCADORIA", "").strip()

# Planilha de entrada
EXCEL_PATH = SCRIPT_DIR / "Emissao_MEL.xlsx"

# Log em CSV
LOG_CSV = SCRIPT_DIR / f"log_emissao_cte_{datetime.now():%Y%m%d_%H%M%S}.csv"

# Colunas de retorno gravadas na planilha
COL_STATUS_AUTOMACAO = "STATUS_AUTOMACAO"
COL_MENSAGEM_AUTOMACAO = "MENSAGEM_AUTOMACAO"
COL_DATA_HORA_AUTOMACAO = "DATA_HORA_AUTOMACAO"

COL_FILA_SEQUENCIA = "FILA_SEQUENCIA"
COL_FILA_OPCAO = "FILA_OPCAO"
COL_FILA_DATA_HORA = "FILA_DATA_HORA"
COL_FILA_USUARIO = "FILA_USUARIO"
COL_FILA_UNIDADE = "FILA_UNIDADE"
COL_FILA_SITUACAO = "FILA_SITUACAO"
COL_FILA_DURACAO = "FILA_DURACAO"

# =====================================================================
# XPATHS - LOGIN
# =====================================================================
XPATH_CAMPO_DOMINIO = "/html/body/form/input[1]"
XPATH_CAMPO_CPF = "/html/body/form/input[2]"
XPATH_CAMPO_USUARIO = "/html/body/form/input[3]"
XPATH_CAMPO_SENHA = "/html/body/form/input[4]"
XPATH_BTN_ENTRAR = "/html/body/form/a"

# =====================================================================
# XPATHS - POPUPS
# =====================================================================
XPATH_BTN_OK_AVISO = "//a[contains(., 'OK')]"
XPATH_ABRIR_OPCAO_156 = "/html/body/div[3]/a[1]"

# =====================================================================
# XPATHS - MENU PRINCIPAL
# =====================================================================
XPATH_MENU_UNIDADE = "/html/body/form/input[1]"
XPATH_MENU_OPCAO = "/html/body/form/input[2]"

# =====================================================================
# XPATHS - TELA DE EMISSÃO (OPÇÃO 6)
# =====================================================================
XPATH_LINK_ROMANEIO = (
    "//a[contains(., 'Romaneio/Packing List') or " "(contains(., 'Romaneio') and contains(., 'Packing'))]"
)

XPATH_LINK_ROUTE_FACILITY = (
    "//a[contains(., 'Route/Facility') or " "contains(., 'Route / Facility') or " "contains(., 'Route Facility')]"
)

XPATH_INPUT_ID_ROTA = "/html/body/form/input[3]"
XPATH_ID_FACILITY = "/html/body/form/input[4]"

XPATH_PLACA_COLETA = "/html/body/form/input[6]"
XPATH_TIPO_MERCADORIA = "/html/body/form/input[7]"
XPATH_COD_ESPECIE = "/html/body/form/input[9]"

XPATH_CNPJ_EXPEDIDOR = "/html/body/form/input[13]"
XPATH_CNPJ_RECEBEDOR = "/html/body/form/input[15]"

XPATH_BTN_GERAR = "/html/body/form/a[1]"

# =====================================================================
# XPATHS - OPÇÃO 156 (FILA)
# =====================================================================
XPATH_156_SEQUENCIA = "/html/body/form/div[2]/div[2]/table[1]/tbody/tr[2]/td[1]/div"

XPATH_156_OPCAO = "/html/body/form/div[2]/div[2]/table[1]/tbody/tr[2]/td[2]"

XPATH_156_DATA_HORA = "/html/body/form/div[2]/div[2]/table[1]/tbody/tr[2]/td[3]/div"

XPATH_156_USUARIO = "/html/body/form/div[2]/div[2]/table[1]/tbody/tr[2]/td[4]/div"

XPATH_156_UNIDADE = "/html/body/form/div[2]/div[2]/table[1]/tbody/tr[2]/td[5]/div"

XPATH_156_SITUACAO = "/html/body/form/div[2]/div[2]/table[1]/tbody/tr[2]/td[7]"

XPATH_156_DURACAO = "/html/body/form/div[2]/div[2]/table[1]/tbody/tr[2]/td[8]/div"


# =====================================================================
# FUNÇÕES AUXILIARES DE SELENIUM
# =====================================================================
def esperar(driver, xpath, timeout=20):
    return WebDriverWait(driver, timeout).until(EC.presence_of_element_located((By.XPATH, xpath)))


def clicar(driver, xpath, timeout=20):
    elemento = WebDriverWait(driver, timeout).until(EC.element_to_be_clickable((By.XPATH, xpath)))

    elemento.click()
    return elemento


def preencher(driver, xpath, valor, timeout=20):
    elemento = esperar(driver, xpath, timeout)

    try:
        elemento.clear()
    except Exception:
        pass

    texto = "" if valor is None else str(valor)
    elemento.send_keys(texto)

    return elemento


# =====================================================================
# PLANILHA E LOG
# =====================================================================
def carregar_planilha():
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"Planilha de emissão não encontrada.\nEsperado em:\n{EXCEL_PATH}")

    df = pd.read_excel(EXCEL_PATH, dtype=str).fillna("")

    df.columns = [str(coluna).strip().upper() for coluna in df.columns]

    obrigatorias = [
        "TRAVEL ID",
        "SIGLA DE ORIGEM",
        "PLACA CAVALO",
        "CNPJ ORIGEM",
        "CNPJ DESTINO",
        "UNIDADE",
    ]

    faltando = [coluna for coluna in obrigatorias if coluna not in df.columns]

    if faltando:
        raise RuntimeError(f"Colunas faltando na planilha: {faltando}")

    registros = []

    for indice, row in enumerate(df.to_dict(orient="records"), start=2):
        # Ignora linhas vazias da tabela (sem Travel ID)
        if not str(row["TRAVEL ID"]).strip():
            continue

        # Ignora linhas que já foram enviadas à fila 156
        if str(row.get("STATUS_AUTOMACAO", "")).strip() == "ENVIADO_FILA_156":
            continue

        row["_LINHA_EXCEL"] = indice
        registros.append(row)

    return registros


def registrar_log(row, status, mensagem):
    colunas_log = [
        "TRAVEL ID",
        "SIGLA DE ORIGEM",
        "SIGLA DE DESTINO",
        "PLACA CAVALO",
        "UNIDADE",
        "CNPJ ORIGEM",
        "CNPJ DESTINO",
        "_LINHA_EXCEL",
        "STATUS",
        "MENSAGEM",
        "DATA_HORA",
    ]

    dados = {**row, "STATUS": status, "MENSAGEM": mensagem, "DATA_HORA": datetime.now().isoformat(timespec="seconds")}

    novo_arquivo = not LOG_CSV.exists()

    with open(LOG_CSV, "a", newline="", encoding="utf-8-sig") as arquivo:
        writer = csv.DictWriter(arquivo, fieldnames=colunas_log, delimiter=";", extrasaction="ignore", restval="")

        if novo_arquivo:
            writer.writeheader()

        writer.writerow(dados)


def normalizar_texto_msg(texto):
    texto = str(texto or "").lower()

    texto = unicodedata.normalize("NFD", texto)

    texto = "".join(caractere for caractere in texto if unicodedata.category(caractere) != "Mn")

    return texto


def garantir_coluna(ws, nome_coluna):
    nome_coluna = nome_coluna.upper()

    for cell in ws[1]:
        cabecalho = str(cell.value or "").strip().upper()

        if cabecalho == nome_coluna:
            return cell.column

    nova_coluna = ws.max_column + 1

    ws.cell(row=1, column=nova_coluna).value = nome_coluna

    return nova_coluna


def atualizar_retorno_planilha(row, status, mensagem, dados_fila=None):
    try:
        linha_excel = row.get("_LINHA_EXCEL")

        if not linha_excel:
            return

        wb = load_workbook(EXCEL_PATH)
        ws = wb.worksheets[0]

        col_status = garantir_coluna(ws, COL_STATUS_AUTOMACAO)
        col_mensagem = garantir_coluna(ws, COL_MENSAGEM_AUTOMACAO)
        col_data = garantir_coluna(ws, COL_DATA_HORA_AUTOMACAO)
        ws.cell(row=linha_excel, column=col_status).value = status
        ws.cell(row=linha_excel, column=col_mensagem).value = mensagem
        ws.cell(row=linha_excel, column=col_data).value = datetime.now().isoformat(timespec="seconds")

        if dados_fila:
            mapa_colunas = {
                COL_FILA_SEQUENCIA: "FILA_SEQUENCIA",
                COL_FILA_OPCAO: "FILA_OPCAO",
                COL_FILA_DATA_HORA: "FILA_DATA_HORA",
                COL_FILA_USUARIO: "FILA_USUARIO",
                COL_FILA_UNIDADE: "FILA_UNIDADE",
                COL_FILA_SITUACAO: "FILA_SITUACAO",
                COL_FILA_DURACAO: "FILA_DURACAO",
            }

            for nome_coluna, chave in mapa_colunas.items():
                coluna_excel = garantir_coluna(ws, nome_coluna)

                ws.cell(row=linha_excel, column=coluna_excel).value = dados_fila.get(chave, "")

        wb.save(EXCEL_PATH)

        print(f"[ok] Retorno gravado na linha {linha_excel} da planilha.")

    except PermissionError:
        print("[warn] Não foi possível atualizar a planilha. " "Verifique se o arquivo está aberto no Excel.")

    except Exception as erro:
        print(f"[warn] Não consegui atualizar retorno na planilha: {erro}")


# =====================================================================
# SSW - LOGIN E NAVEGAÇÃO
# =====================================================================
def login(driver):
    print("[info] Acessando login SSW...")

    driver.get(SSW_URL)

    preencher(driver, XPATH_CAMPO_DOMINIO, SSW_DOMINIO)
    preencher(driver, XPATH_CAMPO_CPF, SSW_CPF)
    preencher(driver, XPATH_CAMPO_USUARIO, SSW_USER)
    preencher(driver, XPATH_CAMPO_SENHA, SSW_PASS)
    clicar(driver, XPATH_BTN_ENTRAR)

    WebDriverWait(driver, 20).until(EC.url_contains("menu01"))

    print("[ok] Login realizado.")


def fechar_janelas_extras(driver, janela_principal):
    for handle in driver.window_handles[:]:
        if handle != janela_principal:
            driver.switch_to.window(handle)
            driver.close()

    driver.switch_to.window(janela_principal)


def abrir_menu_principal(driver):
    driver.get("https://sistema.ssw.inf.br/bin/menu01")

    WebDriverWait(driver, 20).until(EC.url_contains("menu01"))


def definir_unidade_e_abrir_lista(driver, unidade):
    print(f"[info] Definindo unidade: {unidade}")

    abrir_menu_principal(driver)

    campo_unidade = esperar(driver, XPATH_MENU_UNIDADE)

    campo_unidade.clear()
    campo_unidade.send_keys(str(unidade).strip())

    janelas_antes = driver.window_handles[:]

    campo_opcao = esperar(driver, XPATH_MENU_OPCAO)

    campo_opcao.clear()
    campo_opcao.send_keys("6")
    campo_opcao.send_keys(Keys.ENTER)

    print("[ok] Opção 6 enviada.")

    time.sleep(3)

    try:
        WebDriverWait(driver, 10).until(lambda navegador: len(navegador.window_handles) > len(janelas_antes))

        novas_janelas = [handle for handle in driver.window_handles if handle not in janelas_antes]

        if novas_janelas:
            driver.switch_to.window(novas_janelas[-1])

            print("[ok] Foco alterado para janela 006.")

    except Exception:
        print("[info] Opção 6 abriu na mesma janela.")

    WebDriverWait(driver, 20).until(
        lambda navegador: (
            "ssw1602" in navegador.current_url or "006" in navegador.title or "Geração de CTRCs" in navegador.title
        )
    )

    print("[ok] Tela 006 aberta.")

    clicar(driver, XPATH_LINK_ROMANEIO)

    time.sleep(1)

    WebDriverWait(driver, 20).until(EC.element_to_be_clickable((By.XPATH, XPATH_LINK_ROUTE_FACILITY)))

    print("[ok] Tela Romaneio/Packing List aberta.")


def abrir_tela_emissao(driver):
    print("[info] Abrindo Route/Facility...")

    janelas_antes = driver.window_handles[:]
    clicar(driver, XPATH_LINK_ROUTE_FACILITY)

    time.sleep(1)

    WebDriverWait(driver, 20).until(lambda navegador: len(navegador.window_handles) > len(janelas_antes))

    novas_janelas = [handle for handle in driver.window_handles if handle not in janelas_antes]

    if not novas_janelas:
        raise RuntimeError("A tela de emissão não abriu em uma nova janela.")

    driver.switch_to.window(novas_janelas[-1])

    esperar(driver, XPATH_INPUT_ID_ROTA)

    print("[ok] Tela de emissão aberta.")

    time.sleep(1)


def limpar_campos_emissao(driver):
    campos = [
        XPATH_INPUT_ID_ROTA,
        XPATH_ID_FACILITY,
        XPATH_PLACA_COLETA,
        XPATH_TIPO_MERCADORIA,
        XPATH_COD_ESPECIE,
        XPATH_CNPJ_EXPEDIDOR,
        XPATH_CNPJ_RECEBEDOR,
    ]

    for xpath in campos:
        try:
            elemento = esperar(driver, xpath, timeout=5)

            elemento.clear()

        except Exception:
            pass


# =====================================================================
# FILA 156
# =====================================================================
def abrir_156_e_capturar_primeira_linha(driver):
    janela_emissao = driver.current_window_handle
    janelas_antes = driver.window_handles[:]

    print("[info] Abrindo opção 156 pelo popup...")

    # Botão "Abrir a opção 156" do popup
    clicar(driver, XPATH_ABRIR_OPCAO_156)

    WebDriverWait(driver, 20).until(lambda navegador: len(navegador.window_handles) > len(janelas_antes))

    novas_janelas = [handle for handle in driver.window_handles if handle not in janelas_antes]

    if not novas_janelas:
        raise RuntimeError("A janela da opção 156 não foi aberta.")

    janela_156 = novas_janelas[-1]

    driver.switch_to.window(janela_156)

    WebDriverWait(driver, 20).until(EC.presence_of_element_located((By.XPATH, XPATH_156_SEQUENCIA)))

    # Captura a primeira linha da fila (a mais recente)
    dados_fila = {
        "FILA_SEQUENCIA": esperar(driver, XPATH_156_SEQUENCIA).text.strip(),
        "FILA_OPCAO": esperar(driver, XPATH_156_OPCAO).text.strip(),
        "FILA_DATA_HORA": esperar(driver, XPATH_156_DATA_HORA).text.strip(),
        "FILA_USUARIO": esperar(driver, XPATH_156_USUARIO).text.strip(),
        "FILA_UNIDADE": esperar(driver, XPATH_156_UNIDADE).text.strip(),
        "FILA_SITUACAO": esperar(driver, XPATH_156_SITUACAO).text.strip(),
        "FILA_DURACAO": esperar(driver, XPATH_156_DURACAO).text.strip(),
    }

    print("=" * 60)
    print("[fila 156] Dados capturados:")

    for chave, valor in dados_fila.items():
        print(f"{chave}: {valor}")

    print("=" * 60)

    driver.close()

    driver.switch_to.window(janela_emissao)

    print("[ok] Janela 156 fechada. Foco retornado para emissão.")

    return dados_fila


def fechar_popup_ok_simples(driver):
    try:
        botao_ok = WebDriverWait(driver, 5).until(EC.element_to_be_clickable((By.XPATH, XPATH_BTN_OK_AVISO)))
        botao_ok.click()
        time.sleep(1)

        return True

    except Exception:
        return False


# =====================================================================
# EMISSÃO DE UMA LINHA DA PLANILHA
# =====================================================================
def emitir_cte(driver, row, numero_linha):
    try:
        unidade = row.get("UNIDADE", "")
        travel_id = row.get("TRAVEL ID", "")

        if "NÃO ENCONTRADO" in (unidade, row["CNPJ ORIGEM"], row["CNPJ DESTINO"]):
            raise ValueError("Sigla não encontrada no De-Para")

        print(f"[info] Emitindo linha {numero_linha} " f"| Unidade: {unidade} " f"| Travel ID: {travel_id}")

        limpar_campos_emissao(driver)

        preencher(driver, XPATH_INPUT_ID_ROTA, row["TRAVEL ID"])
        time.sleep(1)
        preencher(driver, XPATH_ID_FACILITY, row["SIGLA DE ORIGEM"])
        time.sleep(1)
        preencher(driver, XPATH_PLACA_COLETA, row["PLACA CAVALO"])
        time.sleep(1)
        preencher(driver, XPATH_TIPO_MERCADORIA, SSW_TIPO_MERCADORIA)
        time.sleep(1)
        preencher(driver, XPATH_COD_ESPECIE, "1")
        time.sleep(1)
        preencher(driver, XPATH_CNPJ_EXPEDIDOR, row["CNPJ ORIGEM"])
        time.sleep(1)
        preencher(driver, XPATH_CNPJ_RECEBEDOR, row["CNPJ DESTINO"])
        time.sleep(1)

        # # ===== TESTE: apague estas 2 linhas para emitir de verdade =====
        # input("TESTE: confira os campos no SSW e aperte ENTER...")
        # return
        # ================================================================

        clicar(driver, XPATH_BTN_GERAR)

        time.sleep(2)

        print(f"[ok] Gerar clicado na linha {numero_linha}")

        # Aguarda o SSW responder com um dos dois avisos conhecidos
        WebDriverWait(driver, 20).until(
            lambda navegador: (
                "solicitação enviada para processamento" in navegador.find_element(By.TAG_NAME, "body").text.lower()
                or "programa já está sendo executado" in navegador.find_element(By.TAG_NAME, "body").text.lower()
            )
        )

        texto_tela = driver.find_element(By.TAG_NAME, "body").text

        texto_normalizado = normalizar_texto_msg(texto_tela)

        # -----------------------------------------------------------------
        # CASO 1: foi enviado para a fila 156
        # -----------------------------------------------------------------
        if "solicitacao enviada para processamento" in texto_normalizado:
            dados_fila = abrir_156_e_capturar_primeira_linha(driver)

            mensagem = (
                f"Enviado à fila 156. "
                f"Sequência: {dados_fila.get('FILA_SEQUENCIA')} "
                f"| Situação: {dados_fila.get('FILA_SITUACAO')}"
            )

            registrar_log(row, "ENVIADO_FILA_156", mensagem)

            atualizar_retorno_planilha(row, "ENVIADO_FILA_156", mensagem, dados_fila)

            return

        # -----------------------------------------------------------------
        # CASO 2: o SSW já está executando o programa
        # -----------------------------------------------------------------
        if "programa ja esta sendo executado" in texto_normalizado:
            fechar_popup_ok_simples(driver)

            mensagem = "Programa já está sendo executado " "para os parâmetros informados."

            registrar_log(row, "BLOQUEADO_SSW", mensagem)

            atualizar_retorno_planilha(row, "BLOQUEADO_SSW", mensagem)

            return

        # -----------------------------------------------------------------
        # CASO 3: retorno não reconhecido
        # -----------------------------------------------------------------
        fechar_popup_ok_simples(driver)

        mensagem = "Retorno do SSW não identificado. " "Verificar manualmente."

        registrar_log(row, "RETORNO_NAO_IDENTIFICADO", mensagem)

        atualizar_retorno_planilha(row, "RETORNO_NAO_IDENTIFICADO", mensagem)

    except Exception as erro:
        mensagem_erro = f"{type(erro).__name__}: {erro}"

        print(f"[erro] Linha {numero_linha}: {mensagem_erro}")

        registrar_log(row, "ERRO", mensagem_erro)

        atualizar_retorno_planilha(row, "ERRO", mensagem_erro)


# =====================================================================
# PROGRAMA PRINCIPAL
# =====================================================================
def main():
    if not all([SSW_URL, SSW_DOMINIO, SSW_CPF, SSW_USER, SSW_PASS]):
        raise RuntimeError("Verifique o arquivo .env. Há dados de login faltando.")

    linhas = carregar_planilha()

    if not linhas:
        raise RuntimeError("A planilha não possui linhas para emissão.")

    print(f"[info] Linhas carregadas: {len(linhas)}")

    chrome_options = Options()
    chrome_options.add_argument("--start-maximized")
    chrome_options.add_argument("--disable-popup-blocking")
    chrome_options.add_experimental_option("detach", True)
    driver = webdriver.Chrome(options=chrome_options)
    login(driver)

    janela_principal = driver.current_window_handle

    # Unidades distintas, na ordem em que aparecem na planilha
    unidades = []

    for row in linhas:
        unidade = str(row.get("UNIDADE", "")).strip().upper()

        if unidade and unidade != "NÃO ENCONTRADO" and unidade not in unidades:
            unidades.append(unidade)

    print(f"[info] Unidades encontradas: {unidades}")

    for unidade in unidades:
        print("=" * 80)

        print(f"[lote] Iniciando unidade: {unidade}")

        fechar_janelas_extras(driver, janela_principal)

        linhas_unidade = [row for row in linhas if str(row.get("UNIDADE", "")).strip().upper() == unidade]

        definir_unidade_e_abrir_lista(driver, unidade)

        abrir_tela_emissao(driver)

        for numero_linha, row in enumerate(linhas_unidade, start=1):
            emitir_cte(driver, row, numero_linha)



        fechar_janelas_extras(driver, janela_principal)

        print(f"[lote] Finalizado unidade: {unidade}")

    driver.quit()

    print(f"[fim] Processo finalizado. Log salvo em:\n{LOG_CSV}")


if __name__ == "__main__":
    try:
        main()

    except Exception as erro:
        print("=" * 80)
        print("[erro fatal]")
        print(f"{type(erro).__name__}: {erro}")
        print("=" * 80)

        input("Pressione ENTER para fechar...")